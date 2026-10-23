# -*- coding: utf-8 -*-
"""把两个 xlsx 导入 Grist（表 id 用 ASCII：ItemNames / Achievements）。"""
import json
import os
import pathlib

import requests
from openpyxl import load_workbook

SERVER = "https://sheet.189-24-78-35.sslip.io"
KEY = os.environ.get("GRIST_API_KEY", os.environ.get("GRIST_KEY", ""))
DOC_ID = os.environ.get("GRIST_DOC_ID", "1JRoy2vJar2kh945PXUbFP")

ITEM_COLUMNS = [
    ("english", "英文名"), ("suggested", "建议中文译名"),
    ("alt1", "备选译名1"), ("alt2", "备选译名2"),
    ("alt3", "备选译名3"), ("alt4", "备选译名4"),
    ("final", "最终译名"), ("reason", "选择依据"), ("category", "分类"),
    ("page", "页面状态"), ("source", "翻译来源"), ("status", "处理状态"),
    ("priority", "优先级"), ("owner", "负责人"), ("confirmed", "确认日期"),
    ("note", "备注"), ("wiki_rev", "Wiki 修订"), ("row_hash", "行哈希"),
    ("sync_status", "同步状态"),
]
ACHIEVEMENT_COLUMNS = [
    ("key", "同步键"), ("category", "分类"), ("title", "成就名"),
    ("text", "描述"), ("directions", "达成条件"), ("tutorial", "攻略"),
    ("references", "参考资料"), ("rewards", "奖励"), ("secret", "Secret"),
    ("hidden", "Hidden"), ("obtainable", "Obtainable"),
    ("badge_id", "Badge ID"), ("old_versions", "旧版本"),
    ("wiki_rev", "Wiki 修订"), ("row_hash", "行哈希"),
    ("sync_status", "同步状态"),
]
MODIFIER_COLUMNS = [
    ("key", "同步键"), ("category", "分类"), ("title", "修饰符名"),
    ("description", "中文描述"), ("description_en", "英文描述"),
    ("requirements", "需求"), ("exclusive", "限定"), ("knobs", "旋钮"),
    ("rift", "Rift"), ("progress", "进度"), ("group", "分组"),
    ("reason", "移除原因"), ("wiki_rev", "Wiki 修订"),
    ("row_hash", "行哈希"), ("sync_status", "同步状态"),
]


def cell(row, index):
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def main() -> int:
    session = requests.Session()
    session.headers.update({
        "Authorization": f"Bearer {KEY}",
        "Content-Type": "application/json",
    })

    def apply(actions):
        resp = session.post(f"{SERVER}/api/docs/{DOC_ID}/apply",
                            json=actions, timeout=180)
        if resp.status_code >= 400:
            print("APPLY FAILED", json.dumps(actions, ensure_ascii=False)[:150],
                  resp.status_code, resp.text[:200])
            return None
        return resp.json()

    def tables():
        resp = session.get(f"{SERVER}/api/docs/{DOC_ID}/tables", timeout=60)
        return [t["id"] for t in (resp.json().get("tables") or [])]

    def add_records(table, records):
        resp = session.post(
            f"{SERVER}/api/docs/{DOC_ID}/tables/{table}/records",
            json={"records": [{"fields": r} for r in records]}, timeout=300,
        )
        ok = resp.status_code < 400
        print(f"  写入 {table}: {len(records)} 条 -> {resp.status_code}")
        if not ok:
            print("   ", resp.text[:300])
        return ok

    # 只重建从 xlsx 导入的两张表；Modifiers 由 modifiers_sync.py 维护。
    for name in ("ItemNames", "Achievements"):
        if name in tables():
            apply([["RemoveTable", name]])
    for table_id, columns in (
        ("ItemNames", ITEM_COLUMNS),
        ("Achievements", ACHIEVEMENT_COLUMNS),
    ):
        apply([["AddTable", table_id, [
            {"id": cid, "label": label, "type": "Text"}
            for cid, label in columns
        ]]])
    if "Modifiers" not in tables():
        apply([["AddTable", "Modifiers", [
            {"id": cid, "label": label, "type": "Text"}
            for cid, label in MODIFIER_COLUMNS
        ]]])
    print("表:", tables())

    item_book = load_workbook(
        pathlib.Path("DOORS_物品中文翻译确认表 .xlsx"), read_only=True,
        data_only=True,
    )
    sheet = item_book["物品翻译"]
    item_rows = []
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index < 5:
            continue
        english = cell(row, 1)
        if not english:
            continue
        item_rows.append({
            "english": english,
            "suggested": cell(row, 2),
            "alt1": cell(row, 3), "alt2": cell(row, 4),
            "alt3": cell(row, 5), "alt4": cell(row, 6),
            "final": cell(row, 7), "reason": cell(row, 8),
            "category": cell(row, 9), "page": cell(row, 10),
            "source": cell(row, 11), "status": cell(row, 12),
            "priority": cell(row, 13), "owner": cell(row, 14),
            "confirmed": cell(row, 15), "note": cell(row, 16),
            "sync_status": "来自腾讯文档导出",
        })
    item_book.close()
    print("物品条目:", len(item_rows))
    add_records("ItemNames", item_rows)

    ach_book = load_workbook(
        pathlib.Path("DOORS_Wiki_Achievements.xlsx"), read_only=True,
        data_only=True,
    )
    sheet = ach_book["成就文本"]
    ach_rows = []
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index == 1:
            continue
        key = cell(row, 0)
        if not key:
            continue
        ach_rows.append({
            cid: cell(row, pos) for pos, (cid, _) in enumerate(
                ACHIEVEMENT_COLUMNS
            )
        })
    ach_book.close()
    print("成就条目:", len(ach_rows))
    add_records("Achievements", ach_rows)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
